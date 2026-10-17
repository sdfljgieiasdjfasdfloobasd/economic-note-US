"""Public-only adapters copied from the local dashboard; maintained independently."""
import calendar,csv,io,json,re,zipfile

from datetime import datetime

from decimal import Decimal

from pathlib import Path

from urllib.parse import urljoin

from lxml import html,etree

import openpyxl

import expansion as e

BEA='https://apps.bea.gov/national/Release/XLS/Survey/Section2All_xls.xlsx'

BEA_DETAIL='https://apps.bea.gov/national/Release/XLS/Underlying/Section2All_xls.xlsx'

PROVIDERS={'for':'Federal Reserve Board','bea':'Bureau of Economic Analysis','retail':'U.S. Census Bureau','g19':'Federal Reserve Board','dsr':'Federal Reserve Board','sloos':'Federal Reserve Board','vehicles':'Bureau of Economic Analysis','carts':'Federal Reserve Bank of Chicago','visa':'Visa','sceall':'Federal Reserve Bank of New York','nydebt':'Federal Reserve Bank of New York','cex':'Bureau of Labor Statistics','eia':'Energy Information Administration'}

def item(prefix,code,label,source,publication,unit,frequency='Monthly',season='NSA',note='',**extra):
 return dict(id=f'{prefix}:{code}',label=label,page='pce-nonprice' if prefix=='bea' else 'consumption',category=publication,provider=PROVIDERS.get(prefix,prefix),source=source,publicationId=prefix+':'+publication,publicationLabel=publication,unit=unit,frequency=frequency,seasonal=season,definition=note,observations=[],**extra)

def store(items):
 return [e.save(s) for s in items if any(r['raw'] is not None for r in s['observations'])]

def period(v):
 if isinstance(v,datetime):return v.strftime('%Y-%m-%d')
 s=str(v).strip()
 if re.fullmatch(r'[A-Za-z]{3}-\d{4}',s):return datetime.strptime(s,'%b-%Y').strftime('%Y-%m-01')
 m=re.fullmatch(r'(\d{4})M(\d{2})',s)
 if m:return f'{m[1]}-{m[2]}-01'
 m=re.fullmatch(r'(\d{4})[Q:](\d)',s)
 if m:return f'{m[1]}-{(int(m[2])-1)*3+1:02}-01'
 if re.fullmatch(r'\d{4}-\d\d-\d\d',s):return s
 if re.fullmatch(r'\d{4}',s) and 1900<=int(s)<=2100:return s+'-01-01'
 return e.dt(v)

def bea_parse(b,url,detail=False):
 w=openpyxl.load_workbook(io.BytesIO(b),read_only=True,data_only=True);out=[]
 for s in w:
  if not s.title.endswith('-M'):continue
  if s.title=='T20804-M':continue # Existing price index workbook is the canonical copy.
  rows=list(s.values);title=str(rows[0][0]);units=str(rows[1][0]);datecols={i:period(v) for i,v in enumerate(rows[7]) if i>=3 and period(v)}
  if not datecols:raise ValueError('BEA date header changed: '+s.title)
  release=datetime.strptime(str(rows[4][0]).replace('Data published ',''),'%B %d, %Y').strftime('%Y-%m-%d')
  kind='名目PCE' if 'Personal Consumption' in title and '.5' in title else '実質PCE' if 'Real Personal Consumption' in title else '所得・支出' if '2.6.' in title else '公式伸び率・寄与度'
  path=[]
  for line,r in enumerate(rows[8:],9):
   if len(r)<3 or not r[1]:continue
   label=str(r[1]);indent=len(label)-len(label.lstrip());name=re.sub(r'\\\d+\\','',label.strip())
   while path and path[-1][0]>=indent:path.pop()
   parents=[v[1] for v in path];path.append((indent,name))
   if not r[2]:continue
   code=str(r[2]);unit=units.strip('[]').split(';')[0]
   if 'Contributions to Percent Change' in title:unit='%' if str(r[0])=='1' else 'ポイント'
   if s.title=='T20600-M':
    n=int(r[0]);unit='%' if n==35 or n>=41 else '千人' if n==40 else 'ドル' if n==38 else 'chained (2017) dollars' if n==39 else 'Millions of chained (2017) dollars' if n in [36,37] else 'Millions of dollars'
   elif 'Percent' in unit:unit='ポイント' if 'Contributions' in title else '%'
   # Prefixing generic children with their official parents disambiguates "Other".
   full=' / '.join(parents+[name])
   x=item('bea',s.title+':'+code,full,url,'BEA / '+title,unit,season='SA',note=units+'。公式表の全行。連鎖実質額は加算せず、公表総額を使用。',officialId=code,sourceTable=s.title,hierarchy=parents,depth=len(parents),sourceRow=line,measure=kind,matchLabel=name)
   x['publicationViews']={x['publicationId']:[dict(sourceOrder=line,depth=len(parents),hierarchy=parents,matchLabel=name,label=full)]}
   x['observations']=[dict(e.obs(d,r[c],status='published',sourcePeriod=rows[7][c]),releaseDate=release) for c,d in datecols.items()]
   out.append(x)
 w.close();return out

def merge_publication_views(target,row):
 for pub,views in row.get('publicationViews',{}).items():
  old=target.setdefault('publicationViews',{}).get(pub,[])
  if not isinstance(old,list):old=[old]
  target['publicationViews'][pub]=list({v['sourceOrder']:v for v in old+views}.values())

def bea():
 # Keep table-specific rounded measures distinct; deduplicate identical observations/code/unit.
 standard=e.get(BEA);allrows=bea_parse(standard,BEA)+bea_parse(e.get(BEA_DETAIL),BEA_DETAIL,True);canonical={};out=[]
 prices=json.loads((e.HERE/'data/pce.json').read_text());price_rows=e.workbook_rows(standard,'T20804-M');pricecodes={str(r[2]):next((s for s in prices['series'] if s['name']==re.sub(r'\\\d+\\','',str(r[1])).strip()),None) for r in price_rows[8:] if len(r)>2 and r[2]}
 aliases={}
 for s in allrows:
  ps=pricecodes.get(s['officialId'])
  if ps and s['unit'].startswith('Index') and 'Price Indexes' in s['publicationLabel']:
   code=s['officialId']
   if code not in aliases:
    a=dict(s,id='bea:price:'+code,alias=dict(dataset='pce',series=ps['id'],season='SA'),status='reused',observations=[e.obs(d+'-01',v) for d,v in zip(prices['dates'],ps['bases']['SA'])]);aliases[code]=a
   merge_publication_views(aliases[code],s)
   aliases[code].setdefault('publications',[]).append(dict(id=s['publicationId'],label=s['publicationLabel']));continue
  key=(s['officialId'],s['unit'],s['frequency'])
  old=canonical.get(key)
  if old and old['observations']==s['observations']:
   merge_publication_views(old,s)
   old.setdefault('publications',[{'id':old['publicationId'],'label':old['publicationLabel']}]).append({'id':s['publicationId'],'label':s['publicationLabel']});continue
  canonical[key]=s;out.append(s)
 allitems=out+list(aliases.values());families={}
 for s in allitems:
  code=s['officialId'];measure={'RC':'名目額','RX':'実質額','RG':'価格指数'}.get(code[-2:])
  if measure and (measure!='価格指数' or s['unit'].startswith('Index')):
   s['measure']=measure;families.setdefault(code[:-2],[]).append(s)
 for members in families.values():
  for s in members:s['relatedSeries']=[dict(id=v['id'],label=v['measure']+' · '+v['sourceTable']) for v in members if v['id']!=s['id']]
 return store(out)+[e.summary(a) for a in aliases.values()]

def fed_parse(b,release,url):
 z=zipfile.ZipFile(io.BytesIO(b));root=etree.fromstring(z.read(next(n for n in z.namelist() if n.lower().endswith('_data.xml'))),etree.XMLParser(resolve_entities=False,no_network=True));structure=etree.fromstring(z.read(next(n for n in z.namelist() if n.lower().endswith('_struct.xml'))),etree.XMLParser(resolve_entities=False,no_network=True))
 codes={c.get('id'):{n.get('value'):''.join(n.itertext()).strip() for n in c if n.get('value')} for c in structure.xpath('//*[local-name()="CodeList"]')};out=[]
 for s in root.xpath('//*[local-name()="Series"]'):
  attrs=dict(s.attrib);annotations={a.xpath('string(./*[local-name()="AnnotationType"])'):a.xpath('string(./*[local-name()="AnnotationText"])') for a in s.xpath('.//*[local-name()="Annotation"]')};label=annotations.get('Short Description') or annotations.get('Long Description') or attrs['SERIES_NAME']
  business=release=='sloos' and not bool(re.search(r'consumer|credit card|auto loan|residential|mortgage|home equity',label,re.I))
  freq=codes.get('CL_FREQ',{}).get(attrs['FREQ'],'');unit=attrs['UNIT'].replace('_',' ');mult=attrs.get('UNIT_MULT','1');unit=('%' if unit in ('Percent','Percentage') else ('USD' if unit=='Currency' else unit));unit=({'1000000':'million ','1000000000':'billion ','1000':'thousand '}.get(mult,'')+unit)
  if mult not in ['1','1000','1000000','1000000000']:unit+=' (source multiplier '+mult+')'
  season=attrs.get('SA') or ('SA' if 'seasonally adjusted' in label.lower() else '非該当')
  pub={'g19':'FRB / G.19 Consumer Credit','sloos':'FRB / SLOOS 消費者・住宅向け融資','dsr':'FRB / Household Debt Service Ratios','for':'FRB / Financial Obligations・旧DSR（旧手法）'}[release]
  x=item(release,attrs['SERIES_NAME'],label,url,pub,unit,freq,season,note=annotations.get('Long Description',label)+'。SDMXの公表単位・精度を保持。',officialId=attrs['SERIES_NAME'],dimensions=attrs)
  if business or release=='sloos' and re.search(r'commercial real estate|construction and land|nonfarm nonresidential|multifamily',label,re.I):
   x.update(page='production',publicationId='sloos:business',publicationLabel='FRB / SLOOS 企業・商業不動産向け融資',category='FRB / SLOOS 企業・商業不動産向け融資')
  x['observations']=[e.obs(v.get('TIME_PERIOD'),v.get('OBS_VALUE'),v.get('OBS_STATUS','published')) for v in s.xpath('./*[local-name()="Obs"]')]
  if release=='for':x['definition']+=' 旧手法の履歴。新DSRとは接続しません。FORの公表終了後は補完しません。';x['discontinued']=True
  out.append(x)
 if not out:raise ValueError('SDMX series not found: '+release)
 return out

def fed_zip(release,page):
 # These public full-release endpoints are linked by the official DDP pages.
 # If the chooser is temporarily unavailable, the data endpoint may still work.
 fallback={'g17','g19','for'}
 try:
  t=html.fromstring(e.get(page));links=[urljoin(page,a.get('href')) for a in t.xpath('//a[@href]') if ('.zip' in a.get('href','').lower() or 'filetype=zip' in a.get('href','').lower()) and ('xml' in a.text_content().lower() or 'xml' in a.get('href','').lower())]
  if not links:raise ValueError('公式XML配布リンクが見つかりません: '+release)
  url=links[0]
 except Exception:
  if release not in fallback:raise
  url='https://www.federalreserve.gov/datadownload/Output.aspx?rel='+release.upper()+'&filetype=zip'
 b=e.get(url)
 if not zipfile.is_zipfile(io.BytesIO(b)):raise ValueError('FRB '+release.upper()+' の応答がXML ZIPではありません。保存値を保持します。')
 return b,url

def fed(release):
 page='https://www.federalreserve.gov/'+('data/sloos.htm' if release=='sloos' else 'releases/dsr/' if release=='dsr' else 'datadownload/Download.aspx?rel='+release.upper())
 b,url=fed_zip(release,page)
 return store(fed_parse(b,release,url))

def vehicles():
 url='https://apps.bea.gov/national/xls/gap_hist.xlsx';w=openpyxl.load_workbook(io.BytesIO(e.get(url)),read_only=True,data_only=True);out=[]
 for s in w:
  if not s.title.startswith('Table'):continue
  rows=list(s.values);title=str(rows[1][0]);header=rows[3]
  for c in range(2,len(header)):
   if not header[c]:continue
   label=str(header[c]);season='NSA' if 'Not seasonally' in label else 'SA' if 'Seasonally adjusted' in label else '非該当';unit='百万台・年率' if 'Millions' in label else '千台' if 'Thousands' in label else '季節・営業日調整係数'
   x=item('vehicles',s.title+':'+str(c),title+' / '+label,url,'BEA / Motor Vehicle Unit Retail Sales',unit,season=season,note='BEA公式Motor Vehicle Unit Retail Sales。ドル売上額ではなく販売台数。SAと年率を別系列に保持。')
   for r in rows[5:]:
    if r[0] in calendar.month_name and isinstance(r[1],(int,float)):
     d=f'{int(r[1])}-{list(calendar.month_name).index(r[0]):02}-01'
     x['observations'].append(e.obs(d,r[c],status='projected_factor' if unit=='季節・営業日調整係数' and d>datetime.now().strftime('%Y-%m-%d') else 'published'))
   if unit=='季節・営業日調整係数':x['definition']+=' 将来月の公表調整係数も含みます。販売台数の実績・予測ではありません。'
   out.append(x)
 w.close();return store(out)

def month_name(v):
 s=re.sub(r'[^A-Za-z]','',str(v)).lower()[:3]
 return next((i for i in range(1,13) if calendar.month_abbr[i].lower()==s),None)

def eia():
 import sys
 sys.path.insert(0,str(e.HERE/'vendor'));import xlrd
 out=[]
 for name,url in [('Products Supplied','https://www.eia.gov/dnav/pet/xls/PET_CONS_WPSUP_K_W.xls'),('Gasoline and Diesel Prices','https://www.eia.gov/dnav/pet/xls/PET_PRI_GND_DCUS_NUS_W.xls')]:
  w=xlrd.open_workbook(file_contents=e.get(url))
  for sheet in w.sheets():
   if not sheet.name.startswith('Data'):continue
   codes=sheet.row_values(1);labels=sheet.row_values(2)
   for c in range(1,sheet.ncols):
    code=str(codes[c]);label=str(labels[c]);x=item('eia',code,label,url,'EIA / '+name,'thousand barrels/day' if name=='Products Supplied' else 'USD/gallon','Weekly','NSA','公式週次表の全製品・全価格区分。Products Suppliedは供給側の需要代理指標で、家計消費の直接測定ではありません。',officialId=code)
    x['observations']=[e.obs(xlrd.xldate_as_datetime(sheet.cell_value(i,0),w.datemode).strftime('%Y-%m-%d'),sheet.cell_value(i,c)) for i in range(3,sheet.nrows) if sheet.cell_type(i,0)==xlrd.XL_CELL_DATE];out.append(x)
 return store(out)

def retail_parse(b,url):
 w=openpyxl.load_workbook(io.BytesIO(b),read_only=True,data_only=True);groups={}
 for sheet in w:
  if not sheet.title.isdigit():continue
  rows=list(sheet.values);head=next((r for r in rows[:10] if any(re.match(r'Jan\.? \d{4}',str(v)) for v in r)),None)
  if not head:raise ValueError('MRTS monthly headers changed')
  cols={c:f'{sheet.title}-{month_name(v):02}-01' for c,v in enumerate(head) if c>=2 and month_name(v)};season=None
  for r in rows:
   label=str(r[1] or '').strip()
   if 'NOT ADJUSTED' in label:season='NSA';continue
   if label.startswith('ADJUSTED'):season='SA';continue
   if not season or not label or not any(e.dec(r[c]) for c in cols):continue
   naics=str(r[0] or '').strip();code=naics or {'Retail and food services sales, total':'total','Retail sales, total':'retail','Retail sales and food services excl motor vehicle and parts':'exauto','Retail sales and food services excl gasoline stations':'exgas','Retail sales and food services excl motor vehicle and parts and gasoline stations':'exautogas'}.get(label,label)
   k=(code,season)
   if k not in groups:groups[k]=item('retail',code+':'+season,label,url,'Census / Monthly Retail Trade・全業種','million USD',season=season,note='MRTS公式売上高。物価調整なし。NAICS親項目と内訳を重複加算しません。公表済み全履歴を再取得して改定を反映。',officialId=naics,depth=max(0,len(naics)-3) if naics.isdigit() else 0)
   groups[k]['observations'].extend(e.obs(d,r[c],status='published') for c,d in cols.items())
 w.close();return list(groups.values())
